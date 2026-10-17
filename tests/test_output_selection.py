"""Actual Tk selection, output parity and persistent history regressions."""

import csv
import sqlite3
import subprocess
import sys
from collections.abc import Callable, Iterator
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.core.checkers import CheckerStore
from src.core.check_plan import CHECK_NAMES
from src.core.history_db import HistoryDB, SCHEMA
from src.core.output_settings import DEFAULT_NAMES, OutputSettings
from src.gui.app import MetaInspectorApp
from src.gui.bridge import FinishedEvent
from src.gui.output_actions import OutputActions


@pytest.fixture(scope="session")
def gui_root(tmp_path_factory: pytest.TempPathFactory) -> Iterator[MetaInspectorApp]:
    """Share one interpreter, including repeated collection of this module."""
    app = MetaInspectorApp(profile_dir=tmp_path_factory.mktemp("output_gui") / "profile")
    yield app
    app.destroy()


def reload_output(app: MetaInspectorApp, directory: Path) -> None:
    """Rebuild actual output widgets/settings without recreating the Tk root."""
    for dialog in (app.output_actions.dialog, app.output_actions.history_dialog):
        if dialog is not None and dialog.winfo_exists():
            dialog.destroy()
    toast = getattr(app, "_active_toast", None)
    if toast is not None and toast.winfo_exists():
        toast.destroy()
    app.status_bar.destroy()
    app.profile_dir = directory / "profile"
    app.checker_store = CheckerStore(app.profile_dir)
    app.table.clear()
    app.left_panel.checks_panel.set_selection(CHECK_NAMES)
    app.table.set_checks(CHECK_NAMES)
    app.output_actions = OutputActions(app, app.table, directory)
    app.status_bar = app.output_actions.bar
    app.status_bar.grid(row=2, column=0, columnspan=2, sticky="ew")
    app.update()


@pytest.fixture
def gui_app(gui_root: MetaInspectorApp, tmp_path: Path) -> MetaInspectorApp:
    """Isolate settings and records for each case while keeping the real root."""
    reload_output(gui_root, tmp_path)
    return gui_root


def assert_fresh_app_selection(directory: Path) -> None:
    """Verify saved selection in a new application process with its own Tk root."""
    source = """
import sys
from pathlib import Path
from src.gui.app import MetaInspectorApp
app = MetaInspectorApp(profile_dir=Path(sys.argv[1]) / 'profile')
try:
    assert app.output_actions.names == ('Username', 'Status')
finally:
    app.destroy()
"""
    subprocess.run(
        [sys.executable, "-c", source, str(directory)],
        cwd=Path(__file__).resolve().parents[1], check=True,
        capture_output=True, text=True, timeout=30,
    )


def headers(path: Path) -> list[str]:
    """Read the CSV's actual column headers."""
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return next(csv.reader(handle))


def test_real_popup_drives_clipboard_csv_excel_and_persists(gui_app: MetaInspectorApp, tmp_path: Path,
                                                         monkeypatch: pytest.MonkeyPatch) -> None:
    """Click actual widgets, export actual files and reopen saved preferences."""
    app = gui_app
    destination = tmp_path / "results.csv"
    monkeypatch.setattr("src.gui.actions.filedialog.asksaveasfilename", lambda **kw: str(destination))
    app.update()
    app.table.add_record({"username": "target", "ig_country": "Canada",
                          "ig_date_joined": "April 2016", "data_quality": "Partial"})
    app.status_bar.set_results_available(True)
    app.status_bar.csv_button.invoke()
    before_csv = headers(destination)
    app.status_bar.copy_button.invoke()
    before_tsv = app.clipboard_get().splitlines()[0].split("\t")
    assert before_tsv == list(DEFAULT_NAMES)
    app.status_bar.output_button.invoke()
    app.update()
    popup = app.output_actions.dialog
    assert popup.checkboxes["Username"].cget("state") == "disabled"
    popup.checkboxes["IG Country"].toggle()
    popup.ok_button.invoke()
    app.status_bar.copy_button.invoke()
    assert app.clipboard_get().splitlines()[0].split("\t") == [
        name for name in before_tsv if name != "IG Country"]
    app.status_bar.csv_button.invoke()
    assert headers(destination) == [name for name in before_csv if name != "IG Country"]
    destination = tmp_path / "results.xlsx"
    app.status_bar.xlsx_button.invoke()
    assert [c.value for c in load_workbook(destination).active[1]] == headers(tmp_path / "results.csv")
    assert "IG Country" not in OutputSettings(tmp_path).load()
    assert app.status_bar.copy_button.cget("text").strip().endswith("Copied!")
    assert app._active_toast.winfo_exists()
    app.status_bar.output_button.invoke()
    popup = app.output_actions.dialog
    for name, checkbox in popup.checkboxes.items():
        if name != "Username":
            checkbox.deselect()
    popup.ok_button.invoke()
    app.status_bar.copy_button.invoke()
    assert app.clipboard_get().splitlines()[0] == "Username\tStatus"
    app.status_bar.output_button.invoke()
    app.output_actions.dialog.select_all()
    app.output_actions.dialog.destroy()  # X/Cancel must discard the draft.
    assert app.output_actions.names == ("Username", "Status")
    app.status_bar.output_button.invoke()
    popup = app.output_actions.dialog
    popup.select_all()
    popup.ok_button.invoke()
    app.status_bar.copy_button.invoke()
    selected = app.clipboard_get().splitlines()[0].split("\t")
    assert "IG Country Availability" in selected
    assert "Threads Country Availability" in selected
    assert "Data Quality" in selected
    destination = tmp_path / "selected.csv"
    app.status_bar.csv_button.invoke()
    assert headers(destination) == selected
    app.status_bar.output_button.invoke()
    popup = app.output_actions.dialog
    popup.reset()
    assert tuple(name for name, var in popup.variables.items() if var.get()) == DEFAULT_NAMES
    popup.destroy()
    app.output_actions.apply_selection(("Username", "Status"))
    assert_fresh_app_selection(tmp_path)


@pytest.mark.parametrize("raw", ['{broken', '{"columns":["invalid"]}',
                                      '{"columns":[{}]}', '[]', '\ufffd'])
def test_corrupt_settings_default_and_visible_message(
    gui_app: MetaInspectorApp, tmp_path: Path, raw: str,
) -> None:
    """Corrupt/invalid files never prevent startup and warn in the status bar."""
    (tmp_path / "output_settings.json").write_text(raw, encoding="utf-8")
    settings = OutputSettings(tmp_path)
    assert settings.load() == DEFAULT_NAMES
    assert settings.warning
    app = gui_app
    reload_output(app, tmp_path)
    assert "default columns restored" in app.status_bar._message.cget("text")


def test_history_migrates_and_retains_modern_record(tmp_path: Path) -> None:
    """Read a legacy row and round-trip current metadata through the same DB."""
    path = tmp_path / "history.db"
    with sqlite3.connect(path) as connection:
        connection.executescript(SCHEMA)
        connection.execute("INSERT INTO run_items (run_id, username) VALUES ('old', 'legacy')")
    history = HistoryDB(path)
    assert history.get_run_items("old")[0]["username"] == "legacy"
    record = {"username": "new", "data_quality": "Partial", "ig_status": "private",
              "threads_country_availability": "not_disclosed", "seconds": 3.2,
              "checker_id": "checker_5"}
    run_id = history.save_run([record], "combined", 1, 4)
    assert history.get_run_items(run_id, full_record=True) == [record]


def test_real_history_and_platform_scope(gui_app: MetaInspectorApp, tmp_path: Path) -> None:
    """All field choices reach the engine and finished results open in history."""
    app = gui_app
    for check in CHECK_NAMES:
        app.left_panel.checks_panel.set_selection((check,))
        app.left_panel._checks_changed()
        scraper = app._default_factory(app.left_panel.mode, lambda *args: None)
        assert scraper.checks == (check,)
    record = {"username": "test", "composite_status": "ACTIVE", "data_quality": "Partial"}
    app._total = 1
    app.table.add_record(record)
    app._on_finished(FinishedEvent([record], None, 5.0))
    app.status_bar.history_button.invoke()
    viewer = app.output_actions.history_dialog
    assert viewer.records == [record]
    assert viewer.details.records == [record]


def test_field_cards_forward_scope_and_reject_empty_selection(
    gui_app: MetaInspectorApp, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Actual checkbox edits control the factory and run validation."""
    from src.gui.bridge import RunState

    app = gui_app
    panel = app.left_panel.checks_panel
    try:
        for name, checkbox in panel.checkboxes.items():
            if name not in ("ig_alive", "ig_country"):
                checkbox.toggle()
        assert app.left_panel.requested_checks == ("ig_alive", "ig_country")
        engine = app._default_factory(app.left_panel.mode, lambda *_: None)
        assert engine.mode == "ig_only"
        assert engine.checks == ("ig_alive", "ig_country")
        app.left_panel.apply_state(RunState.RUNNING)
        assert all(
            checkbox.cget("state") == "disabled"
            for checkbox in panel.checkboxes.values()
        )
        app.left_panel.apply_state(RunState.IDLE)
        for checkbox in panel.checkboxes.values():
            checkbox.deselect()
        panel._changed()
        previous = {"username": "previous", "composite_status": "ACTIVE"}
        app.table.set_records([previous])
        app.left_panel.textbox.insert("1.0", "target")
        started = []
        monkeypatch.setattr(
            app.bridge, "start", lambda *args: started.append(args),
        )
        app._on_start()
        assert not started
        assert app.table.records == [previous]
        assert "Select at least one" in app.status_bar._message.cget("text")
        panel.both_button.toggle()
        assert set(panel.selection) == {"ig_alive", "threads_alive"}
        panel.checkboxes["ig_alive"].toggle()
        assert not panel.both_var.get()
    finally:
        app.left_panel.checks_panel.set_selection(CHECK_NAMES)
        app.left_panel._checks_changed()
        app.left_panel.textbox.delete("1.0", "end")


def test_export_presets_draft_cancel_and_filtered_history(
    gui_app: MetaInspectorApp, tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Actual widgets apply field presets and Alive filters."""
    app = gui_app
    records = [
        {"username": "alive", "ig_status": "active",
         "threads_status": "not_found", "ig_country": "Canada",
         "requested_checks": ["ig_alive", "ig_country"]},
        {"username": "unknown", "ig_status": "unknown"},
    ]
    app.table.set_records(records)
    app.status_bar.set_results_available(True)
    app.output_actions.select_output()
    dialog = app.output_actions.dialog
    assert dialog is not None
    dialog.apply_preset("Checked fields")
    chosen = {name for name, var in dialog.variables.items() if var.get()}
    assert chosen == {"Username", "IG Status", "IG Country"}
    dialog.apply_preset("Full export")
    assert all(var.get() for var in dialog.variables.values())
    dialog.apply_preset("Country only")
    dialog.ok_button.invoke()
    assert app.output_actions.names == (
        "Username", "IG Country", "Threads Country",
    )
    app.status_bar.row_filter_var.set("IG Alive only")
    app.status_bar.copy_button.invoke()
    assert app.clipboard_get().splitlines() == [
        "Username\tIG Country\tThreads Country", "alive\tCanada\tN/A",
    ]
    saved = app.output_actions.names
    app.output_actions.select_output()
    app.output_actions.dialog.apply_preset("Age/joined only")
    app.output_actions.dialog.destroy()
    assert app.output_actions.names == saved
    app.output_actions.save_run(records, "ig_only", 2, 1)
    app.output_actions.show_history()
    viewer = app.output_actions.history_dialog
    assert viewer is not None
    viewer.row_filter_var.set("IG Alive only")
    viewer.copy_button.invoke()
    assert len(app.clipboard_get().splitlines()) == 2
    app.status_bar.row_filter_var.set("Both Alive only")
    destination = tmp_path / "should-not-exist.csv"
    monkeypatch.setattr(
        "src.gui.actions.filedialog.asksaveasfilename",
        lambda **_: str(destination),
    )
    app.status_bar.csv_button.invoke()
    assert not destination.exists()
    assert "No completed rows match" in app.status_bar._message.cget("text")
    assert app.table.records == records


def test_live_exports_snapshot_completed_rows_without_stopping_run(
    gui_app: MetaInspectorApp, tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Actual bridge progress permits export before the batch finishes."""
    import threading
    import time

    from src.gui.bridge import RunState, ScraperBridge
    from test_gui_bridge import FakeScraper

    app = gui_app
    original_bridge = app.bridge
    gate = threading.Event()
    gate.set()

    def factory(
        mode: str, callback: Callable[[int, int, dict], None],
    ) -> FakeScraper:
        """Hold the second result while export reads the first final row."""
        def progress(done: int, total: int, record: dict) -> None:
            """Publish one result, then block later work until released."""
            callback(done, total, record)
            gate.clear()
        return FakeScraper(mode, progress, gate=gate)

    app.bridge = ScraperBridge(factory)
    try:
        app.left_panel.textbox.insert("1.0", "first\nsecond")
        app._on_start()
        deadline = time.monotonic() + 4
        while not app.table.records and time.monotonic() < deadline:
            app.update()
            time.sleep(.01)
        assert app.bridge.state is RunState.RUNNING
        assert len(app.table.records) == 1
        app.status_bar.copy_button.invoke()
        assert len(app.clipboard_get().splitlines()) == 2
        path = tmp_path / "live.csv"
        monkeypatch.setattr(
            "src.gui.actions.filedialog.asksaveasfilename",
            lambda **_: str(path),
        )
        app.status_bar.csv_button.invoke()
        with path.open(encoding="utf-8-sig", newline="") as handle:
            assert len(list(csv.reader(handle))) == 2
        path = tmp_path / "live.xlsx"
        app.status_bar.xlsx_button.invoke()
        assert load_workbook(path).active.max_row == 2
        assert app.bridge.state is RunState.RUNNING
        assert len(app.table.records) == 1
    finally:
        gate.set()
        app.bridge.join(3)
        app.bridge.poll()
        app.bridge = original_bridge
        app.left_panel.textbox.delete("1.0", "end")
        app._apply_state()


def test_live_columns_follow_checks_and_retain_recorded_scope(
    gui_app: MetaInspectorApp,
) -> None:
    """Actual Treeview matches Alive boxes and keeps history's own fields."""
    import customtkinter as ctk

    app = gui_app
    panel = app.left_panel.checks_panel
    children = list(app.left_panel.winfo_children())
    widgets = []
    while children:
        widget = children.pop()
        widgets.append(widget)
        children.extend(widget.winfo_children())
    assert not any(isinstance(w, ctk.CTkRadioButton) for w in widgets)
    assert not any(
        "Checking Mode" in str(w.cget("text"))
        for w in widgets if isinstance(w, ctk.CTkLabel)
    )
    for name, checkbox in panel.checkboxes.items():
        if name not in ("ig_alive", "threads_alive"):
            checkbox.toggle()
    expected = (
        "idx", "username", "ig_status", "threads_status",
        "quality", "duration", "error",
    )
    assert tuple(app.table.tree["displaycolumns"]) == expected
    record = {
        "username": "fixture", "composite_status": "ACTIVE",
        "ig_status": "active", "threads_status": "not_found",
        "requested_checks": ["ig_alive", "threads_alive"],
        "data_quality": "Complete",
    }
    app.table.add_record(record)
    item = app.table.tree.get_children()[0]
    assert app.table.tree.set(item, "ig_status") == "ACTIVE"
    assert app.table.tree.set(item, "threads_status") == "NOT FOUND"
    assert app.table._summary.cget("text") == (
        "IG: 1 ACTIVE | Threads: 1 NOT FOUND"
    )
    panel.checkboxes["ig_country"].toggle()
    assert tuple(app.table.tree["displaycolumns"]) == expected
    app.table.set_records([{
        "username": "country", "requested_checks": ["ig_country"],
        "ig_status": "active", "ig_country": "Canada",
    }])
    assert tuple(app.table.tree["displaycolumns"]) == (
        "idx", "username", "ig_country", "quality", "duration", "error",
    )
    item = app.table.tree.get_children()[0]
    assert app.table.tree.set(item, "ig_country") == "Canada"
    app.table.set_records([{"username": "old", "ig_country": "Turkey"}])
    assert tuple(app.table.tree["displaycolumns"]) == ("#all",)
