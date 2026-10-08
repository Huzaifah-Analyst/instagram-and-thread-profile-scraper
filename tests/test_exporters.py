"""Tests for record formatting, the Sheets TSV formatter and CSV/Excel export (TSK-301, TSK-305)."""

import csv
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.core.clipboard import format_tsv
from src.core.export import export_csv, export_xlsx
from src.core.records import EXPORT_COLUMNS, display_status, safe_cell, threads_joined, to_rows

TURKISH = {
    "username": "elif70566", "composite_status": "ACTIVE", "ig_status": "active",
    "ig_country": "Türkiye", "ig_date_joined": "Eylül 2019",
    "threads_status": "active", "threads_country": "Türkiye", "threads_date_joined": "Ekim 2023",
    "threads_badge": "100M+", "seconds": 7.44, "error_message": None,
}
BANNED = {"username": "bad_user", "composite_status": "BANNED", "ig_status": "banned",
          "threads_status": "banned", "seconds": 3.2}
PRIVATE = {"username": "mert23933", "composite_status": "ACTIVE", "ig_status": "private",
           "ig_country": "Iran", "ig_date_joined": "March 2019", "threads_status": "skipped"}


# ---- shared record helpers ------------------------------------------------ #

def test_display_status_private_and_missing() -> None:
    assert display_status(PRIVATE) == "PRIVATE"
    assert display_status(BANNED) == "BANNED"
    assert display_status({}) == "ERROR"


def test_threads_joined_with_and_without_badge() -> None:
    assert threads_joined(TURKISH) == "Ekim 2023 · 100M+"
    assert threads_joined(BANNED) == "N/A"


@pytest.mark.parametrize("raw, expected", [
    ("=HYPERLINK(\"http://x\")", "'=HYPERLINK(\"http://x\")"),
    ("+1 555", "'+1 555"),
    ("-2", "'-2"),
    ("@evil", "'@evil"),
    ("line1\nline2\tx", "line1 line2 x"),
    ("Türkiye", "Türkiye"),
])
def test_safe_cell(raw: str, expected: str) -> None:
    assert safe_cell(raw) == expected


def test_safe_cell_without_guard_keeps_text() -> None:
    assert safe_cell("=1+1", formula_guard=False) == "=1+1"


# ---- TSK-301 Google Sheets TSV --------------------------------------------- #

def test_tsv_header_is_architecture_spec_plus_error() -> None:
    header = format_tsv([]).split("\n")[0]
    assert header == "Username\tStatus\tIG Country\tIG Joined\tThreads Country\tThreads Joined\tError"


def test_tsv_rows_have_seven_cells_each() -> None:
    failed = {**PRIVATE, "error_message": "IG options (...) button not found"}
    lines = format_tsv([TURKISH, BANNED, failed]).split("\n")
    assert len(lines) == 4
    assert all(len(line.split("\t")) == 7 for line in lines)
    assert lines[1] == "elif70566\tACTIVE\tTürkiye\tEylül 2019\tTürkiye\tEkim 2023 · 100M+\t"
    assert lines[2] == "bad_user\tBANNED\tN/A\tN/A\tN/A\tN/A\t"
    assert lines[3].split("\t")[1] == "PRIVATE"
    assert lines[3].endswith("\tIG options (...) button not found")


def test_tsv_value_with_tab_or_newline_cannot_shift_columns() -> None:
    record = {**TURKISH, "ig_country": "Tur\tkey\nX"}
    row = format_tsv([record], include_header=False)
    assert "\n" not in row
    assert len(row.split("\t")) == 7


# ---- TSK-305 CSV / Excel ---------------------------------------------------- #

def test_csv_is_utf8_with_bom_and_round_trips(tmp_path: Path) -> None:
    path = export_csv([TURKISH, BANNED], tmp_path / "out.csv")
    assert path.read_bytes().startswith(b"\xef\xbb\xbf")  # BOM so Excel reads UTF-8
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle))
    assert rows[0] == [name for name, _ in EXPORT_COLUMNS]
    assert rows[1][2] == "Türkiye" and rows[1][3] == "Eylül 2019"
    assert rows[2][:2] == ["bad_user", "BANNED"]


def test_csv_neutralises_formulas(tmp_path: Path) -> None:
    path = export_csv([{**TURKISH, "ig_country": "=cmd|' /C calc'!A0"}], tmp_path / "x.csv")
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle))
    assert rows[1][2].startswith("'=")


def test_xlsx_columns_unicode_and_text_cells(tmp_path: Path) -> None:
    hostile = {**BANNED, "error_message": "=1+1\x07bell"}
    path = export_xlsx([TURKISH, hostile], tmp_path / "out.xlsx")
    sheet = load_workbook(path).active
    rows = [[c.value for c in row] for row in sheet.iter_rows()]
    assert rows[0] == [name for name, _ in EXPORT_COLUMNS]
    assert rows[1][:4] == ["elif70566", "ACTIVE", "Türkiye", "Eylül 2019"]
    error_cell = sheet.cell(row=3, column=rows[0].index("Error") + 1)
    assert error_cell.value == "=1+1bell"  # control char stripped
    assert error_cell.data_type == "s"  # stored as text, never evaluated
    assert sheet.freeze_panes == "A2"


def test_export_rows_include_diagnostics() -> None:
    rows = to_rows([{**TURKISH, "error_message": "IG options (...) button not found"}], EXPORT_COLUMNS)
    header = rows[0]
    assert rows[1][header.index("Error")] == "IG options (...) button not found"
    assert rows[1][header.index("Seconds")] == "7.4"
