"""CSV and Excel export handlers (TSK-305).

CSV is written as UTF-8 with a BOM (``utf-8-sig``): without the BOM, Excel
on Windows opens UTF-8 CSVs in the ANSI codepage and mangles Turkish
characters such as ``Türkiye`` or ``Eylül``.
"""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Union

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter

from src.core.records import EXPORT_COLUMNS, to_rows

PathLike = Union[str, Path]

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="24242A")
_MAX_COLUMN_WIDTH = 60


def export_csv(records: list[dict], path: PathLike) -> Path:
    """Writes records to a CSV file that Excel and Sheets open cleanly.

    Args:
        records: Scraper records.
        path: Destination file.

    Returns:
        The written path.
    """
    target = Path(path)
    with target.open("w", encoding="utf-8-sig", newline="") as handle:
        csv.writer(handle).writerows(to_rows(records, EXPORT_COLUMNS))
    return target


def export_xlsx(records: list[dict], path: PathLike, sheet_title: str = "Results") -> Path:
    """Writes records to an Excel workbook.

    Every cell is stored with an explicit text type, so values that look like
    formulas (``=...``) are shown as text and never evaluated.

    Args:
        records: Scraper records.
        path: Destination ``.xlsx`` file.
        sheet_title: Worksheet name.

    Returns:
        The written path.
    """
    target = Path(path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title

    rows = to_rows(records, EXPORT_COLUMNS, formula_guard=False)
    for row_index, row in enumerate(rows, start=1):
        for col_index, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_index, column=col_index)
            # openpyxl rejects control characters; scraped text may contain them.
            cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)
            cell.data_type = "s"
            if row_index == 1:
                cell.font = _HEADER_FONT
                cell.fill = _HEADER_FILL

    for col_index in range(1, len(rows[0]) + 1):
        width = max(len(row[col_index - 1]) for row in rows) + 2
        sheet.column_dimensions[get_column_letter(col_index)].width = min(width, _MAX_COLUMN_WIDTH)
    sheet.freeze_panes = "A2"

    workbook.save(target)
    return target
